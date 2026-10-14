"""
schedule_controller.py

Weekly Maintenance Schedule: lists the machines that need maintenance on
the upcoming Friday, based on the maintenance matrix workbook.
"""

import calendar
import re
import sys
from datetime import date, timedelta
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
          "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

FRIDAY = 4

DUE_TAG_SHEET = "Due Tag"

# A date counts as "this Friday" when it is within this many days of it.
WINDOW_DAYS = 3

# Last allowable date = scheduled date + days for the frequency.
# Frequencies not listed (W, 2Y, 6Y) use GRACE_DAYS.
LAST_DATE_DAYS = {"M": 4, "3M": 7, "6M": 15, "Y": 30}
GRACE_DAYS = 3

# Order frequencies are written in the output.
FREQUENCY_ORDER = ["W", "M", "3M", "6M", "Y", "2Y", "6Y"]

# Years that must pass since the last 2Y/6Y maintenance before it is due.
LONG_INTERVALS = {"2Y": 2, "6Y": 6}

HEADERS = ["Sl No", "Equipment Name", "Equipment ID", "Frequency",
           "Schedule Date"]


def upcoming_friday(today=None):
    """Today if it is Friday, otherwise the next Friday."""

    today = today or date.today()

    return today + timedelta(days=(FRIDAY - today.weekday()) % 7)


class ScheduleController:

    def __init__(self, excel_file, output_file, friday=None, progress=None):

        self.excel_file = Path(excel_file)
        self.output_file = Path(output_file)
        self.friday = friday or upcoming_friday()
        self.progress = progress
        self.pdf_file = None

    def _report(self, current, total, message):

        if self.progress:
            self.progress(current, total, message)

    # ==================================================
    # Parsing helpers
    # ==================================================

    @staticmethod
    def _tokens(value):

        return {
            token.strip().upper()
            for token in str(value).split(",")
            if token.strip() and token.strip() != "-"
        }

    @staticmethod
    def _days(value):
        """'25, 30' -> (25, 30); '25' -> (25, 25); '-' or '' -> (None, None)."""

        parts = []

        for part in str(value).split(","):

            match = re.search(r"\d+", part)

            if match:
                parts.append(int(match.group()))

        if not parts:
            return None, None

        return parts[0], parts[1] if len(parts) > 1 else parts[0]

    @staticmethod
    def _year(value):

        text = str(value).strip()

        if text in ("", "-", "nan"):
            return None

        try:
            return int(float(text))
        except ValueError:
            parsed = pd.to_datetime(text, errors="coerce")
            return None if pd.isna(parsed) else parsed.year

    def _long_interval_due(self, token, last_done):

        if last_done is None:
            return True

        return self.friday.year - last_done >= LONG_INTERVALS[token]

    # ==================================================
    # Rules
    # ==================================================

    def _candidate_dates(self, day):
        """The matrix day-of-month in the previous, current and next month
        of the Friday, as real dates (clamped for short months)."""

        candidates = []

        for offset in (-1, 0, 1):

            month_index = self.friday.year * 12 + self.friday.month - 1 + offset

            year, month = divmod(month_index, 12)

            month += 1

            last_day = calendar.monthrange(year, month)[1]

            candidates.append(date(year, month, min(day, last_day)))

        return candidates

    def evaluate(self, row):
        """Returns (frequencies, scheduled_dates) due this Friday."""

        frequencies = []

        dates = []

        friday_month = MONTHS[self.friday.month - 1]

        if "W" in self._tokens(row.get(friday_month, "")):

            frequencies.append("W")

            dates.append(self.friday)

        monthly_day, other_day = self._days(row.get("Maintenance Date", ""))

        last_done = self._year(
            row.get("Last 2Y/6Y Done", row.get("Last 2Y Done", ""))
        )

        wanted = {
            "M": monthly_day,
            "3M": other_day, "6M": other_day, "Y": other_day,
            "2Y": other_day, "6Y": other_day,
        }

        for frequency, day in wanted.items():

            if day is None:
                continue

            if frequency in LONG_INTERVALS and not self._long_interval_due(
                    frequency, last_done):
                continue

            for candidate in self._candidate_dates(day):

                if abs((candidate - self.friday).days) > WINDOW_DAYS:
                    continue

                month_tokens = self._tokens(
                    row.get(MONTHS[candidate.month - 1], ""))

                if frequency in month_tokens:

                    frequencies.append(frequency)

                    dates.append(candidate)

                    break

        pairs = sorted(zip(frequencies, dates),
                       key=lambda pair: FREQUENCY_ORDER.index(pair[0]))

        return [f for f, _ in pairs], [d for _, d in pairs]

    # ==================================================
    # Load / build
    # ==================================================

    @staticmethod
    def _sort_key(equipment_id):

        return [int(part) if part.isdigit() else part.lower()
                for part in re.split(r"(\d+)", str(equipment_id).strip())]

    def build_rows(self):

        if not self.excel_file.exists():
            raise FileNotFoundError(
                f"Equipment workbook not found:\n{self.excel_file}")

        self._report(0, 100, "Reading maintenance matrix...")

        workbook = pd.ExcelFile(self.excel_file)

        # The schedule is built from the workbook's "Due Tag" sheet only.
        if DUE_TAG_SHEET not in workbook.sheet_names:
            raise ValueError(
                f"The workbook has no '{DUE_TAG_SHEET}' sheet.\n"
                f"Sheets found: {', '.join(workbook.sheet_names)}")

        sheets = [DUE_TAG_SHEET]

        results = []

        for index, sheet in enumerate(sheets, start=1):

            self._report(index / len(sheets) * 90, 100, f"Checking {sheet}...")

            data = pd.read_excel(workbook, sheet_name=sheet).fillna("")

            data.columns = data.columns.astype(str).str.strip()

            missing = {"Equipment No.", "Maintenance Date"} - set(data.columns)

            if missing:
                raise ValueError(
                    f"Sheet '{sheet}' is missing: {', '.join(sorted(missing))}")

            for _, row in data.iterrows():

                equipment_id = str(row["Equipment No."]).strip()

                if not equipment_id:
                    continue

                frequencies, dates = self.evaluate(row)

                if not frequencies:
                    continue

                scheduled = min(dates)

                # One entry per frequency, used for the due tags.
                items = [
                    {"frequency": frequency, "date": when,
                     "last": when + timedelta(
                         days=LAST_DATE_DAYS.get(frequency, GRACE_DAYS))}
                    for frequency, when in zip(frequencies, dates)
                ]

                results.append({
                    "items": items,
                    "name": str(row.get("Equipment Name", "")).strip(),
                    "id": equipment_id,
                    "frequency": ", ".join(frequencies),
                    "date": scheduled,
                })

        # Same order as the rows in the "Due Tag" sheet.
        return results

    # ==================================================
    # Output
    # ==================================================

    def generate(self):

        rows = self.build_rows()

        self._report(95, 100, "Writing schedule...")

        workbook = Workbook()

        sheet = workbook.active

        sheet.title = "Maintenance Schedule"

        thin = Side(style="thin", color="000000")

        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        sheet.append(HEADERS)

        for number, item in enumerate(rows, start=1):

            sheet.append([number, item["name"], item["id"],
                          item["frequency"], item["date"]])

        widths = [8, 46, 26, 20, 16]

        for column, width in enumerate(widths, start=1):

            sheet.column_dimensions[
                sheet.cell(1, column).column_letter].width = width

        for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row):

            for cell in row:

                cell.border = border

                cell.alignment = Alignment(
                    horizontal="left" if cell.column in (2, 3) else "center",
                    vertical="center", wrap_text=True)

                if cell.row > 1 and cell.column == 5:

                    cell.number_format = "DD-MMM-YYYY"

        for cell in sheet[1]:

            cell.font = Font(bold=True, color="FFFFFF")

            cell.fill = PatternFill("solid", fgColor="1F4E78")

        sheet.freeze_panes = "A2"

        self.output_file.parent.mkdir(parents=True, exist_ok=True)

        workbook.save(self.output_file)

        self.pdf_file = self.output_file.with_suffix(".pdf")

        self._report(98, 100, "Creating PDF...")

        self.write_pdf(rows, self.pdf_file)

        self._report(100, 100, "Done")

        return rows

    # ==================================================
    # PDF
    # ==================================================

    def week_title(self):
        """'Schedule for 3rd Week of October,2026' for the Friday's week."""

        week = (self.friday.day - 1) // 7 + 1

        suffix = {1: "st", 2: "nd", 3: "rd"}.get(week, "th")

        return (f"Schedule for {week}{suffix} Week of "
                f"{self.friday.strftime('%B')},{self.friday.year}")

    def write_pdf(self, rows, pdf_file):

        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.pdfgen import canvas as pdf_canvas
        from reportlab.platypus import (Paragraph, SimpleDocTemplate, Spacer,
                                        Table, TableStyle)

        logo_path = Path(getattr(sys, "_MEIPASS", Path(__file__).parent)) /             "assets" / "header_logo.png"

        page_width, page_height = A4

        # Geometry taken from the company header document (twips -> points).
        table_width = 9351 / 20
        split_x = 1780 / 20
        row1_height = 907 / 20
        row2_height = 485 / 20
        left = (page_width - table_width) / 2
        top = page_height - 12.7 * mm

        def draw_header(canvas, _document):

            canvas.saveState()

            bottom = top - row1_height - row2_height

            canvas.setLineWidth(0.6)
            canvas.setStrokeColor(colors.black)
            canvas.rect(left, bottom, table_width, row1_height + row2_height)
            canvas.line(left, top - row1_height,
                        left + table_width, top - row1_height)
            canvas.line(left + split_x, top, left + split_x, bottom)

            if logo_path.exists():
                canvas.drawImage(
                    str(logo_path), left + 14.75, top - 4.25 - 35.7,
                    width=36.07, height=35.69, mask="auto")

            right = left + table_width - 5.4

            canvas.setFillColor(colors.black)
            canvas.setFont("Times-Bold", 16)
            canvas.drawRightString(right, top - 15, "RENATA PLC")
            canvas.setFont("Times-Bold", 12)
            canvas.drawRightString(right, top - 29.5,
                                   "Rajendrapur General Facility")
            canvas.setFont("Times-Roman", 10)
            canvas.drawRightString(
                right, top - 41.5,
                "Noyapara, Bhawal Mirzapur, Rajendrapur, Gazipur-1700, "
                "Bangladesh")

            middle = top - row1_height - row2_height / 2

            canvas.setFont("Times-Bold", 14)
            canvas.drawString(left + 5.4, middle - 5, "Doc. Title:")
            canvas.setFont("Times-Bold", 15)
            canvas.drawCentredString(
                left + split_x + (table_width - split_x) / 2, middle - 5.5,
                "Preventive Maintenance Schedule")

            canvas.restoreState()

        class NumberedCanvas(pdf_canvas.Canvas):
            """Two-pass canvas so the footer can say 'Page X of Y'."""

            def __init__(self, *args, **kwargs):
                super().__init__(*args, **kwargs)
                self._pages = []

            def showPage(self):
                self._pages.append(dict(self.__dict__))
                self._startPage()

            def save(self):
                total = len(self._pages)
                for state in self._pages:
                    self.__dict__.update(state)
                    self.setFont("Times-Bold", 12)
                    self.drawCentredString(
                        page_width / 2, 12.7 * mm,
                        f"Page {self._pageNumber} of {total}")
                    super().showPage()
                super().save()

        cell = ParagraphStyle("cell", fontName="Times-Roman", fontSize=10,
                              leading=12)

        head = ParagraphStyle("head", parent=cell, fontName="Times-Bold",
                              textColor=colors.black, alignment=1)

        centered = ParagraphStyle("centered", parent=cell, alignment=1)

        document = SimpleDocTemplate(
            str(pdf_file), pagesize=A4,
            leftMargin=left, rightMargin=left,
            topMargin=41 * mm, bottomMargin=22 * mm,
            title="Maintenance Schedule", author="RPGF Engineering")

        story = [
            Paragraph(
                self.week_title(),
                ParagraphStyle("sub", parent=cell, fontName="Times-Bold",
                               fontSize=12, leading=15, alignment=1)),
            Spacer(1, 2 * mm),
        ]

        data = [[Paragraph(text, head) for text in HEADERS]]

        for number, item in enumerate(rows, start=1):

            data.append([
                Paragraph(str(number), centered),
                Paragraph(item["name"], cell),
                Paragraph(item["id"], cell),
                Paragraph(item["frequency"], centered),
                Paragraph(item["date"].strftime("%d-%b-%Y"), centered),
            ])

        # Same width as the header table (165 mm).
        table = Table(
            data, repeatRows=1,
            colWidths=[14 * mm, 58 * mm, 36 * mm, 30 * mm, 27 * mm])

        table.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1),
             [colors.white, colors.HexColor("#E6E6E6")]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))

        story.append(table)

        document.build(
            story, onFirstPage=draw_header, onLaterPages=draw_header,
            canvasmaker=NumberedCanvas)
